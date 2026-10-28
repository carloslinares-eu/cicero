"""End-to-end check of local models on this computer: engine, model download, translation.

Does what a user does in the window, without the window: installs Cicero's own copy of
Ollama if needed, downloads a model, translates a generated three-slide presentation
and checks the result. Run it after changing cicero/core/local/, after updating the
pinned Ollama release, and on each operating system Cicero supports:

    python scripts/check_local_model.py                       # TranslateGemma 4B, into Spanish
    python scripts/check_local_model.py hy-mt2-7b --target de
    python scripts/check_local_model.py --managed             # ignore an Ollama already running

The first run downloads the engine (about 1.6 GB) and the model (3–8 GB).
Exits with 0 when the translation looks right, 1 otherwise.
"""

from __future__ import annotations

import argparse
import csv
import logging
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

from cicero.core.job import JobOptions, TranslationJob  # noqa: E402
from cicero.core.local import catalog  # noqa: E402
from cicero.core.local.hardware import detect_gpu  # noqa: E402
from cicero.core.local.runtime import OllamaRuntime  # noqa: E402
from cicero.core.local.translator import LocalTranslator  # noqa: E402


def build_deck(path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "Northwind Analytics: Q3 Business Review"
    slide.placeholders[1].text = "Prepared for the Board of Directors, October 2026"
    slide.notes_slide.notes_text_frame.text = "Thank everyone for coming and mention that the numbers are preliminary."
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Highlights"
    body = slide.placeholders[1].text_frame
    body.text = "Revenue grew 18% year over year, driven by the new Insight Cloud subscriptions."
    body.add_paragraph().text = "We signed three enterprise customers in the healthcare sector."
    paragraph = body.add_paragraph()
    for text, bold in (("Next quarter, our priority is ", False), ("data privacy compliance", True),
                       (" in the European market.", False)):
        run = paragraph.add_run()
        run.text, run.font.bold = text, bold
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Risks and mitigations"
    table = slide.shapes.add_table(2, 2, Inches(0.5), Inches(1.6), Inches(9), Inches(1.5)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Risk", "Mitigation"
    table.cell(1, 0).text = "Delays in the data centre in Frankfurt"
    table.cell(1, 1).text = "Keep the Dublin region as a fallback until March."
    deck.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("model", nargs="?", default=catalog.DEFAULT_KEY, choices=[m.key for m in catalog.CATALOG])
    parser.add_argument("--target", default="es", help="language code to translate into (default: es)")
    parser.add_argument("--managed", action="store_true", help="use Cicero's own Ollama, not one already running")
    parser.add_argument("--keep", action="store_true", help="keep the engine running and print the output folder")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")

    model = catalog.get(args.model)
    runtime = OllamaRuntime(external_url=None) if args.managed else OllamaRuntime()
    print(f"Graphics card: {detect_gpu() or 'none detected (the model runs on the processor)'}")
    started = time.monotonic()
    if not runtime.is_available():
        print("Installing the engine…", flush=True)
        last = [0.0]

        def report(progress) -> None:
            if time.monotonic() - last[0] >= 10:  # a line every 10 s is enough in a log
                last[0] = time.monotonic()
                print(f"  {progress.stage}: {progress.done / 1e6:,.0f} / {progress.total / 1e6:,.0f} MB", flush=True)

        runtime.install(report)
    client = runtime.client()
    print(f"Engine {client.version()} at {client.base_url} ({'Cicero-managed' if runtime.owns_server else 'existing'})"
          f" after {time.monotonic() - started:.0f} s", flush=True)
    try:
        if not client.has_model(model.ollama_name):
            print(f"Downloading {model.name}…", flush=True)
            client.pull(model.ollama_name)
        started = time.monotonic()
        client.load(model.ollama_name)
        print(f"{model.name} loaded in {time.monotonic() - started:.0f} s", flush=True)

        folder = Path(tempfile.mkdtemp(prefix="cicero-check-"))
        source = folder / "review.pptx"
        build_deck(source)
        output = folder / f"review ({args.target}).pptx"
        translator = LocalTranslator(model, client)
        result = TranslationJob(translator, JobOptions(source, output, args.target, None, write_review=True)).run()
        if result.error:
            print(f"FAILED: {result.error}")
            return 1
        print(f"Translated {result.paragraphs} paragraphs in {result.elapsed:.0f} s")
        with open(result.review_path, encoding="utf-8-sig") as file:
            rows = list(csv.DictReader(file))
        for row in rows:
            print(f"  {row['Original']}\n    -> {row['Translation']}")

        problems = [row["Original"] for row in rows if not row["Translation"].strip()
                    or (row["Translation"] == row["Original"] and any(c.isalpha() for c in row["Original"]))]
        runs = Presentation(str(output)).slides[1].placeholders[1].text_frame.paragraphs[2].runs
        if [run.font.bold for run in runs] != [False, True, False]:
            problems.append(f"formatting: {[(run.text, run.font.bold) for run in runs]}")
        problems += result.warnings
        if problems:
            print("PROBLEMS:\n  " + "\n  ".join(problems))
            return 1
        print(f"OK: {model.name} translated everything and kept the formatting. Output: {output}")
        return 0
    finally:
        if not args.keep:
            runtime.stop()


if __name__ == "__main__":
    raise SystemExit(main())
