from __future__ import annotations

import csv
from pathlib import Path

from pptx import Presentation

from cicero.core.job import JobOptions, Phase, TranslationJob


def all_text(path: Path) -> list[str]:
    deck = Presentation(str(path))
    texts = []
    for slide in deck.slides:
        for shape in slide.shapes:
            texts.extend(_shape_text(shape))
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
    return texts


def _shape_text(shape) -> list[str]:
    if shape.shape_type == 6:  # group
        return [text for child in shape.shapes for text in _shape_text(child)]
    if shape.has_text_frame:
        return [paragraph.text for paragraph in shape.text_frame.paragraphs]
    if getattr(shape, "has_table", False):
        return [cell.text for cell in shape.table.iter_cells()]
    return []


def run_job(translator, source: Path, **options):
    output = source.with_name("out.pptx")
    job = TranslationJob(translator, JobOptions(source, output, "es", **options))
    return job, job.run(), output


def test_translates_every_kind_of_text(fake_translator, sample_deck):
    job, result, output = run_job(fake_translator, sample_deck)
    assert result.succeeded, result.error
    assert job.progress.phase is Phase.FINISHED
    texts = all_text(output)
    for expected in ["QUARTERLY RESULTS", "R&D <DRAFT>", "HELLO BOLD WORLD", "NAME:\t\tCARLOS", "INSIDE GROUP",
                     "NESTED GROUP", "PRODUCT", "PRICE", "APPLE", "SPEAK SLOWLY"]:
        assert expected in texts
    assert "2024 – 42 %" in texts and "3.50" in texts  # nothing to translate, left alone and not sent
    sent = [segment for request in fake_translator.requests for segment in request[0]]
    assert not any("2024" in segment or "3.50" in segment for segment in sent)
    assert result.pages == 3 and result.paragraphs == 10


def test_keeps_run_formatting(fake_translator, sample_deck):
    _, _, output = run_job(fake_translator, sample_deck)
    box = Presentation(str(output)).slides[1].shapes[0]
    runs = box.text_frame.paragraphs[0].runs
    assert [(run.text, run.font.bold) for run in runs] == [("HELLO ", False), ("BOLD", True), (" WORLD", False)]
    assert runs[1].font.size.pt == 18


def test_does_not_add_notes_to_slides_without_them(fake_translator, sample_deck):
    _, _, output = run_job(fake_translator, sample_deck)
    assert [slide.has_notes_slide for slide in Presentation(str(output)).slides] == [True, False, False]


def test_source_language_is_passed_on(fake_translator, sample_deck):
    run_job(fake_translator, sample_deck, source_language="en")
    assert {request[1:] for request in fake_translator.requests} == {("es", "en")}


def test_review_spreadsheet(fake_translator, sample_deck):
    _, result, output = run_job(fake_translator, sample_deck, write_review=True)
    assert result.review_path == output.with_suffix(".csv")
    assert result.review_path is not None
    with open(result.review_path, encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    assert rows[0] == {"Slide": "1", "Location": "Title 1", "Original": "Quarterly results",
                       "Translation": "QUARTERLY RESULTS"}
    assert {"Slide": "1", "Location": "Notes", "Original": "Speak slowly", "Translation": "SPEAK SLOWLY"} in rows


def test_stop_saves_nothing(fake_translator, sample_deck):
    fake_translator.max_segments = 1  # one paragraph per request
    job = TranslationJob(fake_translator, JobOptions(sample_deck, sample_deck.with_name("out.pptx"), "es"))
    def stop_at_the_second_request(number: int) -> None:
        if number == 1:
            job.cancel()

    fake_translator.on_request = stop_at_the_second_request
    result = job.run()
    assert result.cancelled and not result.succeeded
    assert len(fake_translator.requests) == 2
    assert not sample_deck.with_name("out.pptx").exists()


def test_errors_are_reported_not_raised(fake_translator, tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_text("not a presentation")
    _, result, output = run_job(fake_translator, bad)
    assert result.error is not None and "isn't a PowerPoint" in result.error
    assert not output.exists()
    assert list(tmp_path.glob("*.cicero-partial")) == []


def test_progress_counts_paragraphs_slides_and_time(fake_translator, sample_deck):
    fake_translator.max_segments = 1  # one paragraph per request
    seen = []
    job = TranslationJob(fake_translator, JobOptions(sample_deck, sample_deck.with_name("out.pptx"), "es"))
    fake_translator.on_request = lambda number: seen.append(job.progress)
    job.run()
    assert [p.paragraphs_done for p in seen] == list(range(10))
    assert all(p.paragraphs_total == 10 and p.pages_total == 3 for p in seen)
    assert [p.page for p in seen[1:]] == [1, 1, 1, 2, 2, 2, 2, 3, 3]  # slide of the paragraph just translated


def test_remaining_time_follows_the_speed_so_far():
    from cicero.core.job import Progress

    progress = Progress(Phase.TRANSLATING, characters_done=250, characters_total=1000, translating_for=10)
    assert progress.remaining == 30
    assert Progress(Phase.SAVING, characters_done=1000, characters_total=1000, translating_for=10).remaining is None


def smartart_texts(path: Path) -> dict[str, list[str]]:
    from lxml import etree

    from cicero.core.presentation import SMARTART_RELATIONSHIPS

    slide_part = Presentation(str(path)).slides[0].part
    texts = {}
    for relationship in slide_part.rels.values():
        if relationship.reltype in SMARTART_RELATIONSHIPS:
            root = etree.fromstring(relationship.target_part.blob)
            # noinspection HttpUrlsUsage
            texts[relationship.target_part.partname.rsplit("/", 1)[-1]] = [
                node.text for node in root.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t")]
    return texts


def test_translates_smartart_data_and_cached_drawing(fake_translator, smartart_deck):
    _, result, output = run_job(fake_translator, smartart_deck)
    assert result.succeeded, result.error
    assert smartart_texts(output) == {"data1.xml": ["PLAN", "BUILD", "SHIP & LEARN"],
                                      "drawing1.xml": ["PLAN", "BUILD", "SHIP & LEARN"]}
    # The drawing repeats the data's text, so each paragraph is only sent (and billed) once.
    sent = [segment for request in fake_translator.requests for segment in request[0]]
    assert len(sent) == 3 and result.paragraphs == 6


def test_smartart_formatting_is_kept(fake_translator, smartart_deck):
    _, _, output = run_job(fake_translator, smartart_deck)
    slide_part = Presentation(str(output)).slides[0].part
    data = next(r.target_part for r in slide_part.rels.values() if r.target_part.partname.endswith("data1.xml"))
    assert b'b="1"' in data.blob


def master_texts(path: Path) -> list[str]:
    from cicero.core.presentation import iter_text_frames

    deck = Presentation(str(path))
    containers = [master for master in deck.slide_masters]
    containers += [layout for master in deck.slide_masters for layout in master.slide_layouts]
    return [frame.text for container in containers for _, frame in iter_text_frames(container)]


def test_masters_and_layouts_only_when_asked(fake_translator, sample_deck):
    import copy

    from pptx.util import Inches

    deck = Presentation(str(sample_deck))
    box = deck.slides[1].shapes.add_textbox(Inches(0), Inches(7), Inches(3), Inches(0.5))
    box.text_frame.text = "Confidential"
    deck.slide_masters[0].shapes._spTree.append(copy.deepcopy(box._element))  # masters have no add_textbox
    box._element.getparent().remove(box._element)
    deck.save(sample_deck)

    _, plain, output = run_job(fake_translator, sample_deck)
    assert "Confidential" in master_texts(output)
    assert plain.paragraphs == 10

    _, result, output = run_job(fake_translator, sample_deck, include_masters=True)
    texts = master_texts(output)
    assert "CONFIDENTIAL" in texts and "CLICK TO EDIT MASTER TITLE STYLE" in texts
    assert "Confidential" not in texts
    assert result.paragraphs > 10


def test_master_blocks_are_labelled(sample_deck):
    from cicero.core.presentation import PptxDocument

    blocks = [block for block in PptxDocument(sample_deck, include_masters=True).blocks if block.page == 0]
    assert blocks and blocks[0].location.startswith("Master: ")
    assert any(block.location.startswith("Layout: Title Slide › ") for block in blocks)
