from __future__ import annotations

import html
import re
from collections.abc import Sequence
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from cicero.core.translator import Language

LANGUAGES = [Language("en", "English"), Language("es", "Spanish"), Language("zh-CN", "Chinese (Simplified)"),
             Language("pt-PT", "Portuguese (Portugal)")]


class FakeTranslator:
    """Upper-cases the text between tags, like a translator that keeps markup in place."""

    def __init__(self) -> None:
        self.requests: list[tuple[list[str], str, str | None]] = []
        self.on_request = None  # optional hook called before each request

    def languages(self) -> list[Language]:
        return list(LANGUAGES)

    def translate(self, html_segments: Sequence[str], target: str, source: str | None = None) -> list[str]:
        if self.on_request:
            self.on_request(len(self.requests))
        self.requests.append((list(html_segments), target, source))
        return [re.sub(r">([^<]+)<", lambda m: ">" + html.escape(html.unescape(m.group(1)).upper(), quote=False) + "<",
                       segment) for segment in html_segments]


@pytest.fixture
def fake_translator() -> FakeTranslator:
    return FakeTranslator()


@pytest.fixture
def sample_deck(tmp_path: Path) -> Path:
    """A presentation with text in every kind of place Cicero looks."""
    deck = Presentation()
    blank = deck.slide_layouts[6]

    first = deck.slides.add_slide(deck.slide_layouts[0])
    first.shapes.title.text = "Quarterly results"
    first.placeholders[1].text = "R&D <draft>"
    first.notes_slide.notes_text_frame.text = "Speak slowly"

    second = deck.slides.add_slide(blank)
    box = second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    for text, bold in (("Hello ", False), ("bold", True), (" world", False)):
        run = paragraph.add_run()
        run.text = text
        run.font.bold = bold
        run.font.size = Pt(18)
    numbers = box.text_frame.add_paragraph()
    numbers.add_run().text = "2024 – 42 %"
    tabbed = box.text_frame.add_paragraph()
    tabbed.add_run().text = "Name:\t\tCarlos"

    group = second.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1)).text_frame.text = "inside group"
    inner = group.shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1)).text_frame.text = "nested group"

    third = deck.slides.add_slide(blank)
    table = third.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "Product"
    table.cell(0, 1).text = "Price"
    table.cell(1, 0).text = "Apple"
    table.cell(1, 1).text = "3.50"

    path = tmp_path / "deck.pptx"
    deck.save(path)
    return path
