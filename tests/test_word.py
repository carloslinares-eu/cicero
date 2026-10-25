from __future__ import annotations

import csv

import docx

from cicero.core.documents import open_document
from cicero.core.job import JobOptions, TranslationJob


def translate(fake_translator, source, **options):
    output = source.with_name("out.docx")
    result = TranslationJob(fake_translator, JobOptions(source, output, "es", **options)).run()
    assert result.succeeded, result.error
    return result, output


def test_translates_body_tables_headers_and_footers(fake_translator, sample_docx):
    result, output = translate(fake_translator, sample_docx)
    document = docx.Document(str(output))
    body = [paragraph.text for paragraph in document.paragraphs]
    assert body[:3] == ["ANNUAL REPORT", "SALES GREW A LOT", "NAME:\tCARLOS\nMADRID"]
    assert "2024 – 42 %" in body
    cell = document.tables[0].cell(0, 0)
    nested = document.tables[0].cell(0, 1).tables[0].cell(0, 0)
    assert (cell.text, nested.text) == ("REGION", "NESTED CELL")
    section = document.sections[0]
    assert section.header.paragraphs[0].text == "CONFIDENTIAL"
    assert section.footer.paragraphs[0].text == "COMPANY NAME"
    assert result.kind == "Word document" and result.unit is None


def test_keeps_run_formatting_and_links(fake_translator, sample_docx):
    _, output = translate(fake_translator, sample_docx)
    document = docx.Document(str(output))
    runs = document.paragraphs[1].runs
    assert [(run.text, run.bold) for run in runs] == [("SALES ", None), ("GREW", True), (" A LOT", None)]
    assert "OUR WEBSITE" in document.paragraphs[4]._p.xml


def test_runs_with_pictures_are_left_alone(fake_translator, sample_docx):
    _, output = translate(fake_translator, sample_docx)
    paragraph = docx.Document(str(output)).paragraphs[5]
    assert "<w:drawing/>" in paragraph._p.xml
    assert "with picture" in paragraph._p.xml
    assert paragraph.runs[0].text == "PAGE "


def test_locations_and_review_without_page_column(fake_translator, sample_docx):
    locations = {block.location for block in open_document(sample_docx).blocks}
    assert locations == {"Body", "Body › Table", "Header", "Footer"}

    result, _ = translate(fake_translator, sample_docx, write_review=True)
    with open(result.review_path, encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file))
    assert rows[0] == ["Location", "Original", "Translation"]
    assert rows[1] == ["Body", "Annual report", "ANNUAL REPORT"]


def test_unsupported_files_are_explained(fake_translator, tmp_path):
    source = tmp_path / "notes.odt"
    source.write_text("x")
    result = TranslationJob(fake_translator, JobOptions(source, tmp_path / "out.odt", "es")).run()
    assert "can't translate .odt files" in result.error
