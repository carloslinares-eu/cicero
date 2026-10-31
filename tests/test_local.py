"""Local models: catalogue, prompts, the Ollama client, the engine installer and the context-aware translator."""

from __future__ import annotations

import hashlib
import io
import sys
import threading
import time
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from cicero.core.errors import OperationCancelled, TranslationError
from cicero.core.job import JobOptions, Phase, TranslationJob
from cicero.core.local import catalog, languages
from cicero.core.local.catalog import PromptStyle
from cicero.core.local.ollama import OllamaClient
from cicero.core.local.prompts import Brief, Context, clean_answer, from_markers, to_markers, translation_messages
from cicero.core.local.runtime import EngineError, OllamaRuntime, model_on_disk
from cicero.core.local.translator import LocalTranslator, detect_language
from cicero.core.segments import decode_runs, encode_runs
from tests.fake_ollama import FakeOllama


@pytest.fixture
def ollama():
    server = FakeOllama()
    yield server
    server.close()


def installed(server: FakeOllama, key: str) -> catalog.LocalModel:
    model = catalog.get(key)
    server.installed.add(model.ollama_name)
    return model


# -- catalogue ----------------------------------------------------------------------------

def test_catalogue_is_consistent():
    assert len({model.key for model in catalog.CATALOG}) == len(catalog.CATALOG)
    for model in catalog.CATALOG:
        assert all(code in languages.NAMES for code in model.language_codes), model.key
        assert model.gpu_gb in (8, 12, 16) and model.download_gb < model.gpu_gb
        assert model.memory_gb < model.gpu_gb
    assert len(languages.TRANSLATEGEMMA) == 55 and len(languages.HY_MT2) == 38
    assert set(languages.HY_MT2) - set(languages.CHINESE_NAMES) == set()
    assert catalog.get("no-such-model").key == catalog.DEFAULT_KEY


def test_detected_codes_map_to_the_model_languages():
    assert languages.from_detected("zh-cn", languages.TRANSLATEGEMMA) == "zh-Hans"
    assert languages.from_detected("pt", languages.TRANSLATEGEMMA) == "pt-BR"
    assert languages.from_detected("tl", languages.HY_MT2) == "fil"
    assert languages.from_detected("xx", languages.HY_MT2) is None
    assert detect_language("Les résultats du trimestre sont excellents et l'équipe est contente.",
                           languages.TRANSLATEGEMMA) == "fr"


# -- markers and prompts --------------------------------------------------------------------

@pytest.mark.parametrize("runs", [["R&D <draft>"], ["Hello ", "bold", " world"], ["Name:\t\tCarlos\nMadrid"]])
def test_markers_round_trip(runs):
    marked = to_markers(encode_runs([run.strip(" ") for run in runs]))
    assert marked.tagged == (len(runs) > 1)
    assert "<span" not in marked.text and "&amp;" not in marked.text
    assert decode_runs(from_markers(marked.text, marked), len(runs)) == [run.strip(" ") for run in runs]


def test_lost_reordered_and_invented_tags_keep_every_word():
    marked = to_markers(encode_runs(["Hello", "bold", "world"]))
    assert decode_runs(from_markers("<s3>Mundo</s3> <s2>negrita</s2> hola", marked), 3) == ["", "negrita hola",
                                                                                          "Mundo"]
    assert decode_runs(from_markers("Hola negrita mundo", marked), 3) == ["Hola negrita mundo", "", ""]
    assert "x" in decode_runs(from_markers("<s1>a</s1><s9>x</s9>", marked), 3)[0]


@pytest.mark.parametrize("answer, expected", [
    ("Here is the translation: Hola", "Hola"),
    ('"Hola mundo"', "Hola mundo"),
    ("<think>pensando</think>\nHola", "Hola"),
    ("```\nHola\n```", "Hola"),
    ("<target>Hola</target>", "Hola"),
    ("Translation:Hola", "Hola"),
])
def test_answers_are_cleaned(answer, expected):
    assert clean_answer(answer, "Hello") == expected


def test_quotes_in_the_source_are_kept():
    assert clean_answer('"Hola"', '"Hello"') == '"Hola"'


def test_translategemma_gets_googles_prompt():
    model = catalog.get("translategemma-4b")
    [message] = translation_messages(model, to_markers(encode_runs(["Hello"])), "es", "en")
    assert message["content"].startswith("You are a professional English (en) to Spanish (es) translator. Your goal")
    assert message["content"].endswith("Please translate the following English text into Spanish:\n\n\nHello")


def test_hy_mt_uses_context_terminology_and_chinese_when_needed():
    model = catalog.get("hy-mt2-7b")
    context = Context([("Our fruit company", "Nuestra empresa de fruta")], Brief(glossary={"Apple": "Manzana"}))
    [message] = translation_messages(model, to_markers(encode_runs(["Apple sales grew"])), "es", None, context)
    text = message["content"]
    assert "Apple translates to Manzana" in text and "[Background Information]\nOur fruit company" in text
    assert text.endswith("[Source Text]\nApple sales grew")
    [chinese] = translation_messages(model, to_markers(encode_runs(["Hello"])), "zh-Hans", "en")
    assert chinese["content"].startswith("将以下文本翻译为中文")


def test_general_models_get_the_brief_glossary_and_earlier_text():
    model = catalog.get("qwen3.5-9b")
    brief = Brief("A fruit company's results.", {"Apple": "Manzana", "Banana": "Plátano"})
    system, user = translation_messages(model, to_markers(encode_runs(["Apple", "sales"])), "es", None,
                                        Context([("Q3", "T3")], brief))
    assert "About the document: A fruit company's results." in system["content"]
    assert "Apple → Manzana" in system["content"] and "Banana" not in system["content"]  # only terms in the text
    assert "Q3 → T3" in system["content"] and "Keep every tag" in system["content"]
    assert user == {"role": "user", "content": "<s1>Apple</s1><s2>sales</s2>"}


@pytest.mark.parametrize("location, page, unit, expected", [
    ("Notes", 3, "slide", "the speaker notes of slide 3 (written for the presenter, not shown to the audience)"),
    ("Title 1", 2, "slide", "the title of slide 2"),
    ("Subtitle 2", 1, "slide", "the subtitle of slide 1"),
    ("Table: Table 3", 4, "slide", "a table cell of slide 4"),
    ("Body › Table", 0, None, "a table cell"),
    ("Header", 0, None, "the page header"),
    ("Layout: Title Slide › Title 1", 0, "slide", "the slide master or a layout (shared by many slides)"),
    ("Content Placeholder 2", 5, "slide", "the body text of slide 5"),
    ("Text", 7, "page", "the body text of page 7"),
])
def test_places_are_described_for_the_model(location, page, unit, expected):
    from cicero.core.local.prompts import describe_place

    assert describe_place(location, page, unit) == expected


# -- Ollama client -------------------------------------------------------------------------

def test_pull_reports_progress_over_all_files(ollama):
    seen = []
    OllamaClient(ollama.url).pull("translategemma:4b", seen.append)
    assert seen[-1].status == "success" and (seen[-1].completed, seen[-1].total) == (1010, 1010)
    assert OllamaClient(ollama.url).has_model("translategemma:4b")


def test_pull_errors_and_cancel(ollama):
    client = OllamaClient(ollama.url)
    with pytest.raises(TranslationError, match="isn't available for download"):
        client.pull("hf.co/missing/model:Q4")
    stop = threading.Event()
    stop.set()
    with pytest.raises(OperationCancelled):
        client.pull("translategemma:12b", cancel=stop)


def test_model_names_are_compared_like_ollama_does(ollama):
    ollama.installed.add("qwen3.5:latest")
    assert OllamaClient(ollama.url).has_model("QWEN3.5")
    assert not OllamaClient("http://127.0.0.1:9").is_up()


# -- runtime ------------------------------------------------------------------------------

def test_models_are_found_on_disk_without_the_server(tmp_path):
    manifests = tmp_path / "manifests"
    (manifests / "registry.ollama.ai" / "library" / "translategemma").mkdir(parents=True)
    (manifests / "registry.ollama.ai" / "library" / "translategemma" / "4b").write_text("{}")
    (manifests / "hf.co" / "tencent" / "Hy-MT2-7B-GGUF").mkdir(parents=True)
    (manifests / "hf.co" / "tencent" / "Hy-MT2-7B-GGUF" / "Q4_K_M").write_text("{}")
    assert model_on_disk("translategemma:4b", tmp_path)
    assert model_on_disk("hf.co/tencent/hy-mt2-7b-gguf:q4_k_m", tmp_path)  # Ollama may change the case
    assert not model_on_disk("translategemma:12b", tmp_path)
    assert not model_on_disk("qwen3.5:9b", tmp_path / "nowhere")


def fake_release(ollama: FakeOllama, asset: str, damaged: bool = False,
                 files: dict[str, bytes] | None = None) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as bundle:
        for name, content in (files or {"ollama.exe": b"MZ fake engine",
                                        "lib/ollama/readme.txt": b"libraries"}).items():
            bundle.writestr(name, content)
    data = buffer.getvalue()
    digest = hashlib.sha256(b"something else" if damaged else data).hexdigest()
    ollama.files[f"/releases/v9.9.9/{asset}"] = data
    listing = ollama.files.get("/releases/v9.9.9/sha256sum.txt", b"abc  ./other.zip\n").decode()
    ollama.files["/releases/v9.9.9/sha256sum.txt"] = f"{listing}{digest}  ./{asset}\n".encode()
    return data


def test_engine_install_downloads_checks_and_unpacks(ollama, tmp_path, monkeypatch):
    monkeypatch.setattr("cicero.core.local.runtime.release_assets", lambda: ["ollama-windows-amd64.zip"])
    monkeypatch.setattr("cicero.core.local.runtime.sys.platform", "win32")
    data = fake_release(ollama, "ollama-windows-amd64.zip")
    runtime = OllamaRuntime(tmp_path, release="v9.9.9", releases_url=f"{ollama.url}/releases", external_url=None)
    downloads = tmp_path / "downloads"
    downloads.mkdir()
    (downloads / "v9.9.9-ollama-windows-amd64.zip.partial").write_bytes(data[:10])  # an interrupted download
    stages = []
    executable = runtime.install(lambda progress: stages.append(progress.stage))
    assert executable == tmp_path / "ollama" / "v9.9.9" / "ollama.exe"
    assert (tmp_path / "ollama" / "v9.9.9" / "lib" / "ollama" / "readme.txt").is_file()
    assert "Checking the download" in stages and not list(downloads.iterdir())  # archive removed afterward
    assert runtime.managed_executable() == executable


def test_amd_add_on_is_unpacked_into_the_same_folder(ollama, tmp_path, monkeypatch):
    assets = ["ollama-windows-amd64.zip", "ollama-windows-amd64-rocm.zip"]
    monkeypatch.setattr("cicero.core.local.runtime.release_assets", lambda: assets)
    monkeypatch.setattr("cicero.core.local.runtime.sys.platform", "win32")
    fake_release(ollama, assets[0])
    fake_release(ollama, assets[1], files={"lib/ollama/rocm/hip.dll": b"rocm"})
    runtime = OllamaRuntime(tmp_path, release="v9.9.9", releases_url=f"{ollama.url}/releases", external_url=None)
    stages = []
    runtime.install(lambda progress: stages.append(progress.stage))
    assert (tmp_path / "ollama" / "v9.9.9" / "lib" / "ollama" / "rocm" / "hip.dll").is_file()
    assert (tmp_path / "ollama" / "v9.9.9" / "ollama.exe").is_file()
    assert "Downloading the local model engine (2 of 2)" in stages


@pytest.mark.parametrize("platform_name, machine, amd, expected", [
    ("win32", "AMD64", False, ["ollama-windows-amd64.zip"]),
    ("win32", "AMD64", True, ["ollama-windows-amd64.zip", "ollama-windows-amd64-rocm.zip"]),
    ("win32", "ARM64", False, ["ollama-windows-arm64.zip"]),
    ("linux", "x86_64", False, ["ollama-linux-amd64.tar.zst"]),
    ("linux", "x86_64", True, ["ollama-linux-amd64.tar.zst", "ollama-linux-amd64-rocm.tar.zst"]),
    ("linux", "aarch64", False, ["ollama-linux-arm64.tar.zst"]),
    ("darwin", "arm64", False, ["ollama-darwin.tgz"]),
    ("freebsd14", "amd64", False, []),
])
def test_release_assets_per_system(monkeypatch, platform_name, machine, amd, expected):
    from cicero.core.local.runtime import release_assets

    monkeypatch.setattr("cicero.core.local.runtime.sys.platform", platform_name)
    monkeypatch.setattr("cicero.core.local.runtime.platform.machine", lambda: machine)
    assert release_assets(amd_gpu=amd) == expected


def test_linux_tar_zst_is_unpacked(tmp_path):
    import tarfile

    zstandard = pytest.importorskip("zstandard")  # installed with Cicero on Linux
    from cicero.core.local.runtime import _extract

    raw = io.BytesIO()
    with tarfile.open(fileobj=raw, mode="w") as bundle:
        data = b"#!/bin/sh\necho fake ollama\n"
        info = tarfile.TarInfo("bin/ollama")
        info.size, info.mode = len(data), 0o755
        bundle.addfile(info, io.BytesIO(data))
    archive = tmp_path / "ollama-linux-amd64.tar.zst"
    archive.write_bytes(zstandard.ZstdCompressor().compress(raw.getvalue()))
    _extract(archive, tmp_path / "out")
    assert (tmp_path / "out" / "bin" / "ollama").read_bytes().startswith(b"#!/bin/sh")


@pytest.mark.parametrize("with_filters", [True, False])
@pytest.mark.filterwarnings("ignore:Python 3.14 will:DeprecationWarning")  # from simulating 3.11.2 on a newer one
def test_tar_archives_unpack_safely_with_and_without_filters(tmp_path, monkeypatch, with_filters):
    import tarfile

    from cicero.core.local.runtime import _extract_tar

    if with_filters and not hasattr(tarfile, "data_filter"):
        pytest.skip("this Python has no extraction filters (e.g. Debian 12's 3.11.2): only the fallback applies")
    monkeypatch.setattr("cicero.core.local.runtime.TAR_FILTERS", with_filters)  # False: as on Debian 12's 3.11.2

    def archive(entries: list[tuple[str, str | None]]) -> io.BytesIO:
        raw = io.BytesIO()
        with tarfile.open(fileobj=raw, mode="w") as writer:
            for name, link in entries:
                info = tarfile.TarInfo(name)
                if link is None:
                    info.size, info.mode = 4, 0o4755
                    writer.addfile(info, io.BytesIO(b"data"))
                else:
                    info.type, info.linkname = tarfile.SYMTYPE, link
                    writer.addfile(info)
        raw.seek(0)
        return raw

    members: list[tuple[str, str | None]] = [("bin/ollama", None), ("lib/ollama/libcudart.so.12.8", None)]
    if sys.platform != "win32":  # creating symlinks needs special rights on Windows; Linux archives have them
        members.append(("lib/ollama/libcudart.so.12", "libcudart.so.12.8"))
    good = archive(members)
    with tarfile.open(fileobj=good, mode="r|") as bundle:
        _extract_tar(bundle, tmp_path / "ok")
    assert (tmp_path / "ok" / "bin" / "ollama").read_bytes() == b"data"
    for bad in ([("../escape", None)], [("lib/link", "../../outside")]):
        with tarfile.open(fileobj=archive(bad), mode="r|") as bundle, pytest.raises(Exception):
            _extract_tar(bundle, tmp_path / "bad")
    assert not (tmp_path / "escape").exists()


def test_amd_memory_is_read_from_sysfs(tmp_path):
    from cicero.core.local.hardware import _amd_cards

    for card, vendor, memory in (("card0", "0x8086", "0"), ("card1", "0x1002", str(16 * 1024**3))):
        device = tmp_path / card / "device"
        device.mkdir(parents=True)
        (device / "vendor").write_text(vendor + "\n")
        (device / "mem_info_vram_total").write_text(memory)
    cards = _amd_cards(tmp_path)
    assert [(card.name, card.memory_gb) for card in cards] == [("AMD Radeon", 16.0)]


def test_the_note_says_where_a_model_runs_and_whether_memory_is_enough():
    from cicero.core.local.hardware import Gpu, fit_note, system_memory_gb

    small, large = catalog.get("translategemma-4b"), catalog.get("qwen3.5-9b")
    assert fit_note(large, Gpu("NVIDIA GeForce RTX 3070 Ti", 8.0), 32.0) == "Fits your RTX 3070 Ti (8 GB)"
    # An old laptop: 8 GB of memory (7.7 usable), and no graphics card or a 2 GB one.
    assert fit_note(small, None, 7.7) == "Runs best with 8 GB of graphics memory, or slowly on the processor"
    assert fit_note(large, None, 7.7).startswith("Too large for this computer's 7.7 GB of memory")
    assert fit_note(large, Gpu("GeForce 940MX", 2.0), 7.7).startswith("Larger than your 2 GB card")
    assert fit_note(large, None, None).startswith("Runs best with 8")  # memory unknown: no verdict
    assert (system_memory_gb() or 0) > 0.5


def test_a_running_ollama_service_is_asked_which_models_it_has(ollama, tmp_path):
    ollama.installed.add("translategemma:4b")
    runtime = OllamaRuntime(tmp_path, external_url=ollama.url)
    assert runtime.model_ready("translategemma:4b") and not runtime.model_ready("qwen3.5:9b")


def test_damaged_engine_download_is_refused(ollama, tmp_path, monkeypatch):
    monkeypatch.setattr("cicero.core.local.runtime.release_assets", lambda: ["ollama-windows-amd64.zip"])
    fake_release(ollama, "ollama-windows-amd64.zip", damaged=True)
    runtime = OllamaRuntime(tmp_path, release="v9.9.9", releases_url=f"{ollama.url}/releases", external_url=None)
    with pytest.raises(EngineError, match="checksum"):
        runtime.install()
    assert not any(path.suffix == ".partial" for path in (tmp_path / "downloads").iterdir())


def test_an_ollama_the_user_runs_is_used_as_is(ollama, tmp_path):
    runtime = OllamaRuntime(tmp_path, external_url=ollama.url)
    assert runtime.is_available() and runtime.client().base_url == ollama.url
    assert not runtime.owns_server
    runtime.stop()  # must not touch someone else's server
    assert OllamaClient(ollama.url).is_up()


FAKE_SERVE = '''
import http.server, json, os

host, port = os.environ["OLLAMA_HOST"].rsplit(":", 1)


class Version(http.server.BaseHTTPRequestHandler):
    def do_GET(self):
        body = json.dumps({"version": "0.0.0"}).encode()
        self.send_response(200)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *args):
        pass


http.server.HTTPServer((host, int(port)), Version).serve_forever()
'''


@pytest.mark.skipif(not sys.platform.startswith("linux"), reason="the parent-death signal is Linux's")
def test_the_engine_outlives_the_thread_that_started_it(tmp_path):
    # The window starts the engine from its Connect or Download thread, which ends right after.
    engine = tmp_path / "ollama" / "v-test" / "ollama"
    engine.parent.mkdir(parents=True)
    engine.write_text(f"#!{sys.executable}\n{FAKE_SERVE}")
    engine.chmod(0o755)
    runtime = OllamaRuntime(tmp_path, release="v-test", external_url=None)
    clients: list[OllamaClient] = []
    worker = threading.Thread(target=lambda: clients.append(runtime.client()))
    worker.start()
    worker.join()
    time.sleep(1)  # the kernel's signal, if any, arrives as the thread ends
    try:
        assert runtime.owns_server and clients[0].is_up()
    finally:
        runtime.stop()
    assert not runtime.owns_server


def test_missing_engine_is_explained(tmp_path, monkeypatch):
    monkeypatch.setattr("cicero.core.local.runtime.shutil.which", lambda _name: None)
    runtime = OllamaRuntime(tmp_path, external_url=None)
    assert not runtime.is_available()
    with pytest.raises(EngineError, match="isn't installed yet"):
        runtime.client()


# -- translating ----------------------------------------------------------------------------

def run_job(model, ollama, deck, **options):
    translator = LocalTranslator(model, OllamaClient(ollama.url))
    output = deck.with_name("out.pptx")
    job = TranslationJob(translator, JobOptions(deck, output, "es", **options))
    result = job.run()
    assert result.succeeded, result.error
    return translator, Presentation(str(output)), job


def test_translategemma_detects_the_source_and_keeps_formatting(ollama, sample_deck):
    model = installed(ollama, "translategemma-4b")
    translator, deck, job = run_job(model, ollama, sample_deck)
    assert translator.detected_source == "en"
    assert all("English (en) to Spanish (es)" in chat["messages"][0]["content"] for chat in ollama.chats)
    runs = deck.slides[1].shapes[0].text_frame.paragraphs[0].runs
    assert [(run.text, run.font.bold) for run in runs] == [("HELLO ", False), ("BOLD", True), (" WORLD", False)]
    assert deck.slides[0].shapes.title.text == "QUARTERLY RESULTS"
    assert all(chat["options"]["temperature"] == 0 for chat in ollama.chats)
    assert job.progress.phase is Phase.FINISHED


def test_general_model_reads_the_document_first(ollama, sample_deck):
    model = installed(ollama, "qwen3.5-9b")
    translator, deck, _ = run_job(model, ollama, sample_deck)
    first = ollama.chats[0]
    assert first["format"]["required"] == ["summary", "glossary"] and "Quarterly results" in \
        first["messages"][0]["content"]
    assert translator.brief is not None and translator.brief.glossary == {"Apple": "Manzana"}
    apple = next(chat for chat in ollama.chats[1:] if chat["messages"][-1]["content"] == "Apple")
    assert "Apple → Manzana" in apple["messages"][0]["content"]
    assert all(chat.get("think") is False for chat in ollama.chats)
    notes = next(chat for chat in ollama.chats if chat["messages"][-1]["content"] == "Speak slowly")
    assert "is from the speaker notes of slide 1" in notes["messages"][0]["content"]
    # The second paragraph of slide 1 is sent with the first one and its translation as context.
    subtitle = next(chat for chat in ollama.chats if chat["messages"][-1]["content"] == "R&D <draft>")
    assert "Quarterly results → QUARTERLY RESULTS" in subtitle["messages"][0]["content"]


def test_hy_mt_gets_earlier_paragraphs_as_background(ollama, sample_deck):
    model = installed(ollama, "hy-mt2-7b")
    run_job(model, ollama, sample_deck, source_language="en")
    subtitle = next(chat for chat in ollama.chats if chat["messages"][0]["content"].endswith("R&D <draft>"))
    assert "[Background Information]\nQuarterly results" in subtitle["messages"][0]["content"]
    assert ollama.chats[0]["options"]["top_k"] == 20


def test_rambling_answers_are_retried_and_empty_ones_kept_in_the_original(ollama, sample_deck):
    model = installed(ollama, "translategemma-4b")
    calls = []

    def ramble_once(request):
        calls.append(request)
        text = request["messages"][-1]["content"].rsplit("\n\n", 1)[1]
        if text == "Speak slowly":
            return ""
        if len(calls) == 1:
            return "Sure! " + "This is a long explanation. " * 20
        return text.upper()

    ollama.answer_override = ramble_once
    translator, deck, _ = run_job(model, ollama, sample_deck)
    assert deck.slides[0].shapes.title.text == "QUARTERLY RESULTS"  # the retry
    assert deck.slides[0].notes_slide.notes_text_frame.text == "Speak slowly"
    assert translator.warnings == [
        "1 paragraph(s) came back empty from TranslateGemma 4B and were left in the original language."]


def test_a_model_that_is_not_downloaded_is_explained(ollama, sample_deck):
    translator = LocalTranslator(catalog.get("translategemma-12b"), OllamaClient(ollama.url))
    result = TranslationJob(translator, JobOptions(sample_deck, sample_deck.with_name("o.pptx"), "es", "en")).run()
    assert result.error is not None and "isn't downloaded" in result.error


def test_every_style_has_a_prompt():
    marked = to_markers(encode_runs(["Hi"]))
    for style in PromptStyle:
        model = next(model for model in catalog.CATALOG if model.style is style)
        assert translation_messages(model, marked, "es", "en")[-1]["content"].endswith("Hi")


def test_the_cloud_style_of_detecting_language_isnt_needed_for_context_models(ollama, sample_deck):
    model = installed(ollama, "hy-mt2-7b")
    translator, _, _ = run_job(model, ollama, sample_deck)
    assert translator.detected_source is None  # Hy-MT works it out itself


def test_language_list_comes_from_the_catalogue():
    translator = LocalTranslator(catalog.get("hy-mt2-7b"), OllamaClient("http://127.0.0.1:9"))
    names = [language.name for language in translator.languages()]
    assert len(names) == 38 and "Cantonese" in names and names == sorted(names, key=str.casefold)


def test_model_label_and_paths():
    model = catalog.get("translategemma-4b")
    assert model.label() == "TranslateGemma 4B · 3.3 GB · 8 GB GPU"
    assert isinstance(Path(model.homepage).name, str)
