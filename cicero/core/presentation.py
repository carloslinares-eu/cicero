"""Finding every piece of translatable text in a presentation and putting translations back."""

from __future__ import annotations

import os
import tempfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path

import pptx
import pptx.exc
from pptx.presentation import Presentation
from pptx.shapes.group import GroupShape
from pptx.text.text import TextFrame, _Run

from cicero.core.errors import PresentationError
from cicero.core.segments import decode_runs, encode_runs, needs_translation, split_edges


@dataclass(slots=True)
class TextBlock:
    """One paragraph to translate: its runs that contain text, and the HTML sent to the translator."""

    slide: int  # 1-based slide number
    location: str  # shape name, "Table: …" or "Notes"
    runs: list[_Run]
    edges: list[tuple[str, str]]  # whitespace stripped from the start and end of each run
    original: str  # the paragraph as plain text
    source_html: str
    translation: str = field(default="")  # plain text, filled in by ``apply``


def open_presentation(path: str | Path) -> Presentation:
    try:
        return pptx.Presentation(str(path))
    except FileNotFoundError:
        raise PresentationError(f"File not found:\n{path}") from None
    except PermissionError:
        raise PresentationError(f"Can't open the file. Is it open in another program?\n{path}") from None
    except (pptx.exc.PackageNotFoundError, KeyError, ValueError) as exc:
        raise PresentationError(
            f"This isn't a PowerPoint .pptx file, or it is damaged.\n{path}\n\n"
            "Older .ppt files must first be saved as .pptx from PowerPoint.\n\n"
            f"Details: {exc}"
        ) from None


def collect_blocks(presentation: Presentation) -> list[TextBlock]:
    """Every paragraph with letters in it, in reading order: slide by slide, shapes, tables, then notes."""
    blocks = []
    for number, location, frame in iter_text_frames(presentation):
        for paragraph in frame.paragraphs:
            runs = [run for run in paragraph.runs if run.text.strip()]
            if not runs or not needs_translation("".join(run.text for run in runs)):
                continue
            pieces = [split_edges(run.text) for run in runs]
            blocks.append(TextBlock(
                slide=number,
                location=location,
                runs=runs,
                edges=[(before, after) for before, _, after in pieces],
                original="".join(run.text for run in paragraph.runs),
                source_html=encode_runs([core for _, core, _ in pieces]),
            ))
    return blocks


def iter_text_frames(presentation: Presentation) -> Iterator[tuple[int, str, TextFrame]]:
    for number, slide in enumerate(presentation.slides, start=1):
        for shape in _iter_shapes(slide.shapes):
            if shape.has_text_frame:
                yield number, shape.name, shape.text_frame
            elif getattr(shape, "has_table", False):
                for cell in shape.table.iter_cells():
                    if not cell.is_spanned:
                        yield number, f"Table: {shape.name}", cell.text_frame
        if slide.has_notes_slide:  # reading slide.notes_slide would create an empty one
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                yield number, "Notes", frame


def _iter_shapes(shapes) -> Iterator:
    """Shapes on a slide, looking inside groups (and groups of groups)."""
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def apply(block: TextBlock, translated_html: str) -> None:
    """Write a translation into the block's runs. Each run keeps its own formatting."""
    texts = decode_runs(translated_html, len(block.runs))
    for run, (before, after), text in zip(block.runs, block.edges, texts):
        run.text = f"{before}{text}{after}" if text else ""
    block.translation = "".join(run.text for run in block.runs)


def save_presentation(presentation: Presentation, path: str | Path) -> None:
    """Save through a temporary file, so an existing file is only replaced by a complete one."""
    path = Path(path)
    descriptor, temporary = tempfile.mkstemp(prefix=f".{path.stem}.", suffix=".cicero-partial", dir=path.parent)
    os.close(descriptor)
    try:
        presentation.save(temporary)
        os.replace(temporary, path)
    except PermissionError:
        raise PresentationError(f"Can't write the file. Is it open in PowerPoint?\n{path}") from None
    except OSError as exc:
        raise PresentationError(f"Can't save the file:\n{path}\n\n{exc}") from None
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)
